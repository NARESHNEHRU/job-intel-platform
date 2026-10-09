import io
import re

from backend.app.services.normalize import SKILLS

AMBIGUOUS = {"rest", "node", "spring"}


def _build_patterns():
    patterns = []
    for canon, variants in SKILLS.items():
        for v in [canon] + variants:
            if len(v) <= 3 or v in AMBIGUOUS:
                forms, flags = {canon, v.upper()}, 0      # must be written like "SQL" / "REST"
            else:
                forms, flags = {v}, re.IGNORECASE
            for f in forms:
                patterns.append((canon, re.compile(
                    r"(?<![A-Za-z0-9+#])" + re.escape(f) + r"(?![A-Za-z0-9+#])", flags)))
    return patterns


PATTERNS = _build_patterns()


def extract_skills(text: str) -> list[str]:
    """Known skills found in the text, in the order they first appear."""
    first_seen = {}
    for canon, pat in PATTERNS:
        m = pat.search(text)
        if m and (canon not in first_seen or m.start() < first_seen[canon]):
            first_seen[canon] = m.start()
    return [c for c, _ in sorted(first_seen.items(), key=lambda x: x[1])]


def read_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Please upload a PDF, DOCX or TXT file")